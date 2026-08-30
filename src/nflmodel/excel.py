"""
excel.py — builds the NFL_Betting_Model workbook.

Visual language deliberately matches the MLB workbook: navy headers, yellow
input cells, gray calculated cells, green/red conditional formatting on edges.
Structure differs because NFL is a weekly sport -- the primary sheet is a slate
of ~16 games ranked by edge, not a single-game form.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .config import STAKING
from .data import TEAMS
from .market import format_spread

# ── Colors (same palette as the MLB model) ──
NAVY = "1B2A4A"
WHITE = "FFFFFF"
GREEN = "2E7D32"
LIGHT_YELLOW = "FFFDE7"
LIGHT_GRAY = "F5F5F5"
LIGHT_GREEN = "C8E6C9"
LIGHT_RED = "FFCDD2"
LIGHT_ORANGE = "FFE0B2"
GRAY_BLUE = "607D8B"
DARK_GOLD = "F9A825"
ROW_TINT = "EEF2F7"


def fill(hex_color):
    return PatternFill(fill_type="solid", fgColor=hex_color)


def font(bold=False, color=WHITE, size=11, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic)


def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


def apply_header(ws, ref, text, merge_to=None, size=16):
    c = ws[ref]
    c.value, c.fill, c.font, c.alignment = text, fill(NAVY), font(True, WHITE, size), center()
    if merge_to:
        ws.merge_cells(f"{ref}:{merge_to}")


def apply_section(ws, ref, text, merge_to=None):
    c = ws[ref]
    c.value, c.fill, c.font, c.alignment = text, fill(NAVY), font(True, WHITE, 11), left()
    if merge_to:
        ws.merge_cells(f"{ref}:{merge_to}")


def set_widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _header_row(ws, row, labels, start_col=1):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=label)
        c.fill, c.font, c.alignment, c.border = fill(NAVY), font(True, WHITE, 11), center(), thin_border()
    ws.row_dimensions[row].height = 24


# ─────────────────────────────────────────────
# BUILD
# ─────────────────────────────────────────────

def read_existing_bets(path: str) -> list[list]:
    """
    Pull manually-entered Bet Tracker rows out of a workbook we are about to
    overwrite.

    The weekly runner is meant to be run twice -- once when lines open and
    again before kickoff. Rebuilding the workbook from scratch on the second
    run silently destroyed everything logged after the first, which is the
    worst possible failure for a bet log: quiet, total, and only noticed later.
    """
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path)
        if "Bet Tracker" not in wb.sheetnames:
            return []
        ws = wb["Bet Tracker"]
    except Exception:
        return []

    rows = []
    for r in range(12, 501):
        # A..I are user-entered; J/K/N/O are formulas we rewrite each run.
        vals = [ws.cell(row=r, column=c).value for c in range(1, 10)]
        extras = [ws.cell(row=r, column=c).value for c in (12, 13, 16)]
        if any(v not in (None, "") for v in vals):
            rows.append(vals + extras)
    return rows


def build_workbook(
    slate: pd.DataFrame,
    ratings: pd.DataFrame,
    season: int,
    week: int,
    backtest_summary: dict | None = None,
    path: str = "NFL_Betting_Model.xlsx",
    pts_table: list | None = None,
) -> str:
    preserved = read_existing_bets(path)

    wb = Workbook()

    _sheet_lists(wb, pts_table)
    _sheet_slate(wb, slate, season, week)
    _sheet_detail(wb, slate, season, week)
    _sheet_ratings(wb, ratings, season, week)
    _sheet_tracker(wb, preserved)
    _sheet_reference(wb, backtest_summary)

    wb.save(path)
    return path


def _sheet_lists(wb, pts_table=None):
    ws = wb.active
    ws.title = "Lists"
    ws.sheet_state = "hidden"
    for i, t in enumerate(TEAMS, start=1):
        ws.cell(row=i, column=1, value=t)
    for i, v in enumerate(["W", "L", "Push"], start=1):
        ws.cell(row=i, column=2, value=v)

    # Columns D:E — win-probability value of one point of spread, by line.
    # Used by the tracker to convert points of CLV into probability. This is a
    # lookup rather than a constant because the point crossing 3 is worth four
    # times a point at 2.
    for i, (line, per_pt) in enumerate(pts_table or [], start=1):
        ws.cell(row=i, column=4, value=line)
        ws.cell(row=i, column=5, value=per_pt)


def _sheet_slate(wb, slate, season, week):
    ws = wb.create_sheet("Weekly Slate")
    ws.tab_color = NAVY
    set_widths(ws, {"A": 6, "B": 22, "C": 11, "D": 11, "E": 10, "F": 10, "G": 10,
                    "H": 10, "I": 11, "J": 22, "K": 10, "L": 11})

    apply_header(ws, "A1", f"NFL BETTING MODEL — {season} WEEK {week}", merge_to="L1", size=16)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 8

    cols = ["#", "Matchup", "Model Line", "Vegas Line", "Edge (pts)",
            "Cover %", "Push %", "Win %", "Conf", "Recommendation", "Odds", "Stake"]
    _header_row(ws, 3, cols)
    ws.freeze_panes = "A4"

    for i, (_, g) in enumerate(slate.iterrows(), start=1):
        r = 3 + i
        vals = [
            i,
            f"{g['away_team']} @ {g['home_team']}",
            format_spread(g["home_team"], g["fair_spread"]),
            format_spread(g["home_team"], g["spread_line"]) if pd.notna(g["spread_line"]) else "—",
            g["spread_edge_pts"],
            g["home_cover_prob"] if g["spread_edge_pts"] > 0 else g["away_cover_prob"],
            g["push_prob"],
            g["home_win_prob"],
            g["confidence"],
            g["recommendation"],
            g["bet_odds"] if g["bet_odds"] else "",
            g["stake"] if g["stake"] else "",
        ]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.border, c.alignment = thin_border(), center() if j != 2 and j != 10 else left()
            c.font = Font(color="000000", size=11)
            if i % 2 == 0:
                c.fill = fill(ROW_TINT)

        ws.cell(row=r, column=5).number_format = "+0.0;-0.0"
        ws.cell(row=r, column=5).font = Font(bold=True, size=11)
        for col in (6, 7, 8):
            ws.cell(row=r, column=col).number_format = "0.0%"
        ws.cell(row=r, column=12).number_format = '"$"#,##0'
        ws.cell(row=r, column=10).font = Font(bold=True, size=11)

    last = 3 + len(slate)
    if last > 3:
        rng = f"E4:E{last}"
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"ABS(E4)>=3"], fill=fill(LIGHT_GREEN), font=Font(bold=True, color="1B5E20")))
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"AND(ABS(E4)>=1.5,ABS(E4)<3)"], fill=fill(LIGHT_ORANGE), font=Font(bold=True, color="E65100")))
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"ABS(E4)<1.5"], fill=fill(LIGHT_GRAY), font=Font(color="616161")))

        rec = f"J4:J{last}"
        ws.conditional_formatting.add(rec, FormulaRule(
            formula=['LEFT(J4,3)="BET"'], fill=fill(LIGHT_GREEN), font=Font(bold=True, color="1B5E20")))
        ws.conditional_formatting.add(rec, FormulaRule(
            formula=['J4="NO BET"'], fill=fill(LIGHT_GRAY), font=Font(color="616161")))

    n = last + 2
    apply_section(ws, f"A{n}", "SLATE SUMMARY", merge_to=f"L{n}")
    bets = slate[slate["stake"] > 0]
    for k, (label, val) in enumerate([
        ("Games on slate", len(slate)),
        ("Qualifying bets", len(bets)),
        ("Total staked", f"${bets['stake'].sum():,.0f}"),
        ("% of bankroll", f"{bets['stake'].sum()/STAKING.bankroll*100:.1f}%"),
        ("Largest edge", f"{slate['spread_edge_pts'].abs().max():.1f} pts"),
    ], start=1):
        ws.cell(row=n + k, column=1, value=label).font = Font(bold=True, size=11)
        ws.merge_cells(start_row=n + k, start_column=1, end_row=n + k, end_column=3)
        c = ws.cell(row=n + k, column=4, value=val)
        c.font, c.alignment = Font(bold=True, color=NAVY, size=11), center()

    note = n + 7
    ws.cell(row=note, column=1,
            value="→ Model Line is what the model makes the game. Edge is the disagreement with Vegas, in points. "
                  "See Reference & Glossary for the backtest and what these numbers are worth.")
    ws.cell(row=note, column=1).font = Font(italic=True, color=GREEN, size=10)
    ws.merge_cells(start_row=note, start_column=1, end_row=note, end_column=12)


def _sheet_detail(wb, slate, season, week):
    ws = wb.create_sheet("Game Detail")
    ws.tab_color = GRAY_BLUE
    set_widths(ws, {"A": 28, "B": 18, "C": 18, "D": 4, "E": 28, "F": 18})

    apply_header(ws, "A1", f"GAME DETAIL — {season} WEEK {week}", merge_to="F1", size=16)
    ws.row_dimensions[1].height = 30

    row = 3
    for _, g in slate.iterrows():
        apply_section(ws, f"A{row}", f"{g['away_team']} @ {g['home_team']}", merge_to="F" + str(row))
        row += 1

        pairs = [
            ("Model projected margin", f"{g['projected_margin']:+.2f}", "Vegas line",
             format_spread(g["home_team"], g["spread_line"]) if pd.notna(g["spread_line"]) else "—"),
            ("Model line", format_spread(g["home_team"], g["fair_spread"]), "Edge (points)", f"{g['spread_edge_pts']:+.2f}"),
            ("Home cover %", f"{g['home_cover_prob']:.1%}", "Away cover %", f"{g['away_cover_prob']:.1%}"),
            ("Push %", f"{g['push_prob']:.1%}", "Home win %", f"{g['home_win_prob']:.1%}"),
            ("Model fair ML (home)", f"{g['home_ml_fair']:+.0f}", "Actual ML (home)", f"{g['home_ml']:+.0f}" if pd.notna(g["home_ml"]) else "—"),
            ("Model fair ML (away)", f"{g['away_ml_fair']:+.0f}", "Actual ML (away)", f"{g['away_ml']:+.0f}" if pd.notna(g["away_ml"]) else "—"),
            ("ML edge home (de-vigged)", f"{g['ml_edge_home']:+.1%}" if pd.notna(g["ml_edge_home"]) else "—",
             "ML edge away (de-vigged)", f"{g['ml_edge_away']:+.1%}" if pd.notna(g["ml_edge_away"]) else "—"),
            ("Situational adjustment", f"{g['situational_adj']:+.2f}", "Confidence", g["confidence"]),
            ("RECOMMENDATION", g["recommendation"], "Stake", f"${g['stake']:,.0f}" if g["stake"] else "—"),
        ]
        for lab_l, val_l, lab_r, val_r in pairs:
            for col, (lab, val) in ((1, (lab_l, val_l)), (5, (lab_r, val_r))):
                lc = ws.cell(row=row, column=col, value=lab)
                lc.font, lc.fill, lc.border = Font(bold=(lab == "RECOMMENDATION"), color="000000", size=11), fill("E8EDF3"), thin_border()
                vc = ws.cell(row=row, column=col + 1, value=val)
                vc.font, vc.alignment, vc.border = Font(bold=True, color=NAVY, size=11), center(), thin_border()
                vc.fill = fill(LIGHT_GRAY)
            row += 1
        row += 1


def _sheet_ratings(wb, ratings, season, week):
    ws = wb.create_sheet("Power Ratings")
    ws.tab_color = GREEN
    set_widths(ws, {"A": 8, "B": 14, "C": 16, "D": 40})

    apply_header(ws, "A1", f"POWER RATINGS — entering {season} Week {week}", merge_to="D1", size=16)
    ws.row_dimensions[1].height = 30
    _header_row(ws, 3, ["Rank", "Team", "Rating (pts)", "Interpretation"])
    ws.freeze_panes = "A4"

    for i, (_, r) in enumerate(ratings.iterrows(), start=1):
        row = 3 + i
        vals = [int(r["rank"]), r["team"], round(float(r["rating"]), 2),
                f"{abs(r['rating']):.1f} pts {'better' if r['rating'] >= 0 else 'worse'} than an average team on a neutral field"]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=j, value=v)
            c.border = thin_border()
            c.alignment = center() if j != 4 else left()
            c.font = Font(color="000000", size=11, bold=(j == 3))
            if i % 2 == 0:
                c.fill = fill(ROW_TINT)
        ws.cell(row=row, column=3).number_format = "+0.00;-0.00"

    n = 3 + len(ratings) + 2
    ws.cell(row=n, column=1, value="Ratings are opponent-adjusted and include the projected starting quarterback. "
                                   "A matchup's model line = home rating − away rating + home field advantage.")
    ws.cell(row=n, column=1).font = Font(italic=True, color=GREEN, size=10)
    ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=4)


def _sheet_tracker(wb, preserved=None):
    ws = wb.create_sheet("Bet Tracker")
    ws.tab_color = DARK_GOLD
    set_widths(ws, {"A": 11, "B": 6, "C": 22, "D": 12, "E": 18, "F": 11, "G": 9,
                    "H": 9, "I": 8, "J": 12, "K": 13, "L": 12, "M": 12,
                    "N": 10, "O": 10, "P": 11})

    apply_header(ws, "A1", "BET TRACKER", merge_to="P1", size=16)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 8

    apply_section(ws, "A3", "SUMMARY", merge_to="F3")
    summary = [
        (4, "Starting Bankroll", STAKING.bankroll, '"$"#,##0.00'),
        (5, "Total Bets", "=COUNTA($A$12:$A$500)", None),
        (6, "Record (W-L-P)", '=COUNTIF($I$12:$I$500,"W")&"-"&COUNTIF($I$12:$I$500,"L")&"-"&COUNTIF($I$12:$I$500,"Push")', None),
        (7, "Win %", '=IFERROR(COUNTIF($I$12:$I$500,"W")/(COUNTIF($I$12:$I$500,"W")+COUNTIF($I$12:$I$500,"L")),0)', "0.0%"),
        (8, "Total P&L", "=SUM($J$12:$J$500)", '"$"#,##0.00'),
        (9, "ROI", "=IFERROR(SUM($J$12:$J$500)/SUM($H$12:$H$500),0)", "0.0%"),
        (10, "Current Bankroll", "=$B$4+SUM($J$12:$J$500)", '"$"#,##0.00'),
    ]
    for row, label, formula, fmt in summary:
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = Font(bold=True, color="000000", size=11)
        c = ws.cell(row=row, column=2, value=formula)
        c.font, c.alignment, c.border = Font(bold=True, color=NAVY, size=12), center(), thin_border()
        if fmt:
            c.number_format = fmt
        if row == 4:
            c.fill = fill(LIGHT_YELLOW)

    # CLV summary — the metric that tells you if you're sharp before P&L can.
    ws.cell(row=4, column=4, value="Avg CLV").font = Font(bold=True, size=11)
    clv = ws.cell(row=4, column=5, value='=IFERROR(AVERAGE($O$12:$O$500),"")')
    clv.font, clv.alignment, clv.number_format, clv.border = Font(bold=True, color=NAVY, size=12), center(), "0.00%", thin_border()
    ws.cell(row=5, column=4, value="Beat close %").font = Font(bold=True, size=11)
    bc = ws.cell(row=5, column=5, value='=IFERROR(COUNTIF($O$12:$O$500,">0")/COUNT($O$12:$O$500),"")')
    bc.font, bc.alignment, bc.number_format, bc.border = Font(bold=True, color=NAVY, size=12), center(), "0.0%", thin_border()
    ws.cell(row=4, column=6, value="Avg CLV (pts)").font = Font(bold=True, size=11)
    cp = ws.cell(row=4, column=7, value='=IFERROR(AVERAGE($N$12:$N$500),"")')
    cp.font, cp.alignment, cp.number_format, cp.border = Font(bold=True, color=NAVY, size=12), center(), "+0.00;-0.00", thin_border()
    ws.cell(row=6, column=4, value="↑ Positive CLV is the earliest real evidence of edge. For spreads, points is the honest unit.").font = Font(italic=True, color=GREEN, size=10)
    ws.merge_cells("D6:F6")

    headers = ["Date", "Week", "Matchup", "Market", "Bet Side", "Line Taken",
               "Odds", "Stake", "Result", "P&L", "Bankroll",
               "Closing Line", "Closing Odds", "CLV (pts)", "CLV (%)",
               "Model Edge"]
    _header_row(ws, 11, headers)
    ws.freeze_panes = "A12"

    dv = DataValidation(type="list", formula1='"W,L,Push"', allow_blank=True)
    dv.sqref = "I12:I500"
    ws.add_data_validation(dv)

    dv_mkt = DataValidation(type="list", formula1='"SPREAD,MONEYLINE"', allow_blank=True)
    dv_mkt.sqref = "D12:D500"
    ws.add_data_validation(dv_mkt)

    for row in range(12, 501):
        # J — P&L. A push returns the stake, so it is zero, not a loss.
        ws.cell(row=row, column=10, value=(
            f'=IF($I{row}="W",IF($G{row}<0,$H{row}*100/ABS($G{row}),$H{row}*$G{row}/100),'
            f'IF($I{row}="L",-$H{row},IF($I{row}="Push",0,"")))'
        )).number_format = '"$"#,##0.00'

        # K — running bankroll
        ws.cell(row=row, column=11, value=(
            f'=IF(COUNTA($A$12:$A{row})=0,"",$B$4+SUM($J$12:$J{row}))'
        )).number_format = '"$"#,##0.00'

        # N — CLV in POINTS. Both numbers are quoted from your side, so
        # taken-minus-close is signed correctly: +3.5 taken vs +1.5 close = +2.
        ws.cell(row=row, column=14, value=(
            f'=IF(OR($D{row}<>"SPREAD",$F{row}="",$L{row}=""),"",$F{row}-$L{row})'
        )).number_format = "+0.0;-0.0"

        # O — CLV in probability. Moneyline uses the two prices. Spread
        # converts points via the per-line table on Lists!D:E, because the
        # point that crosses 3 is worth four times a point at 2.
        ws.cell(row=row, column=15, value=(
            f'=IF($D{row}="MONEYLINE",'
            f'IF(OR($G{row}="",$M{row}=""),"",'
            f'IF($M{row}<0,-$M{row}/(-$M{row}+100),100/($M{row}+100))'
            f'-IF($G{row}<0,-$G{row}/(-$G{row}+100),100/($G{row}+100))),'
            f'IF($N{row}="","",'
            f'$N{row}*IFERROR(VLOOKUP(ABS($F{row}),Lists!$D:$E,2,FALSE),0.03)))'
        )).number_format = "0.00%"

        for col in (10, 11, 14, 15):
            c = ws.cell(row=row, column=col)
            c.font, c.alignment, c.border, c.fill = (
                Font(color="000000", size=11), center(), thin_border(), fill(LIGHT_GRAY))
        for col in list(range(1, 10)) + [12, 13, 16]:
            ws.cell(row=row, column=col).border = thin_border()

    for rng in ("N12:N500", "O12:O500"):
        col = rng[0]
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'AND({col}12<>"",{col}12>0)'], fill=fill(LIGHT_GREEN),
            font=Font(bold=True, color="1B5E20")))
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'AND({col}12<>"",{col}12<0)'], fill=fill(LIGHT_RED),
            font=Font(bold=True, color="B71C1C")))

    # Restore anything the user had already logged.
    for i, rec in enumerate(preserved or []):
        r = 12 + i
        for c, v in enumerate(rec[:9], start=1):      # A..I user-entered
            ws.cell(row=r, column=c, value=v)
        for off, col in enumerate((12, 13, 16)):      # closing line/odds, model edge
            idx = 9 + off
            if len(rec) > idx and rec[idx] is not None:
                ws.cell(row=r, column=col, value=rec[idx])


def _sheet_reference(wb, backtest_summary):
    ws = wb.create_sheet("Reference & Glossary")
    ws.tab_color = GRAY_BLUE
    set_widths(ws, {"A": 30, "B": 88})

    apply_header(ws, "A1", "REFERENCE & GLOSSARY", merge_to="B1", size=16)
    ws.row_dimensions[1].height = 30

    row = 3

    def section(title):
        nonlocal row
        apply_section(ws, f"A{row}", title, merge_to="B" + str(row))
        row += 1

    def entry(term, text):
        nonlocal row
        c = ws.cell(row=row, column=1, value=term)
        c.font, c.fill, c.border, c.alignment = Font(bold=True, color="000000", size=11), fill("E8EDF3"), thin_border(), left()
        d = ws.cell(row=row, column=2, value=text)
        d.font, d.border, d.alignment = Font(color="000000", size=11), thin_border(), left()
        ws.row_dimensions[row].height = max(16, 15 * (len(text) // 95 + 1))
        row += 1

    section("HOW THE MODEL WORKS")
    entry("Power rating", "Each team's strength in points versus an average team on a neutral field, fit by ridge regression across every game so it is opponent-adjusted.")
    entry("Regression target", "A blend of actual scoring margin and an EPA-implied margin. EPA predicts the future better than past points do because it strips out turnover luck and garbage time.")
    entry("Quarterback", "Team and QB are fit jointly, so a rating reflects who is actually starting. About 15% of team-games are started by a non-primary QB, and that swing is worth ~6 points.")
    entry("Model line", "home rating − away rating + home field advantage. This is the spread the model would set.")
    entry("Edge (points)", "Model line minus the Vegas line. The model's disagreement with the market.")
    entry("Home field", "Fitted from data at ~2.1 points, not the folk-wisdom 3. It was 0.2 in the fanless 2020 season, which is good evidence the effect is crowd-driven.")

    section("MARKET TERMS")
    entry("Vig / juice", "The book's built-in margin. Two sides at -110 imply 104.8% total probability; the 4.8% excess is the hold.")
    entry("De-vigging", "Stripping the vig to recover the market's true opinion. Comparing your model to raw implied odds overstates or understates your edge on every bet — this model always de-vigs first.")
    entry("Push", "The margin lands exactly on the spread and the stake is returned. Real money: whole-number spreads push about 2.7% of the time, and ignoring that overstates edge.")
    entry("Closing line value (CLV)", "Whether the number you bet was better than the number at kickoff. The earliest reliable signal that a model has genuine edge — it shows up in weeks, where profit takes seasons.")
    entry("Break-even at -110", "52.38%. Anything below that loses money no matter how it feels.")

    section("STAKING")
    entry("Kelly criterion", "Stake sized to the edge and the odds. Full Kelly maximizes long-run growth but swings violently.")
    entry("Half Kelly", f"This model stakes at {STAKING.kelly_fraction:.0%} of Kelly, capped at {STAKING.max_bet_pct:.1%} of bankroll — roughly three-quarters of the growth at half the volatility.")

    section("WHAT THE BACKTEST FOUND")
    if backtest_summary:
        entry("Method", "Walk-forward: to predict week W the model only ever sees games finished before week W, and is refit from scratch each week. No result the model is scored on informs it.")
        entry("Games scored", f"{backtest_summary.get('n_games', 0):,} games, {backtest_summary.get('seasons','')}")
        entry("Model error (MAE)", f"{backtest_summary.get('model_mae', 0):.2f} points per game")
        entry("Closing line error (MAE)", f"{backtest_summary.get('market_mae', 0):.2f} points per game")
        entry("Verdict", backtest_summary.get("verdict", ""))
    entry("Situational factors", "Rest, weather, travel, divisional games and Week 17-18 spots were each measured against 5,431 games of closing lines. None showed mispricing at even |t|=2, so all ship at zero. They are not missing — they were tested and rejected.")

    section("HONEST LIMITS")
    entry("The line is very good", "The closing spread is among the most accurate forecasts in any domain. Most public models do not beat it, and a model that claims a large edge is usually leaking future information.")
    entry("Sample size", "A 53% ATS model needs hundreds of bets before its win rate is distinguishable from noise. Judge this by CLV first and P&L much later.")
    entry("Bet what you can lose", "Sizing rules only control the rate of ruin. They do not remove it.")
