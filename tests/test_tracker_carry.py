"""
The Bet Tracker is carried into a new workbook as the sheet itself.

Jameson, 2026-09-30: the tracker "should only be for you to observe not to
touch and meddle with". excel._start_workbook loads his newest workbook and
drops every other sheet; the model's sheets are rebuilt around it. These
tests pin that nothing on his sheet changes -- values, formulas, colours,
dropdowns, merged cells, a note typed past the last column -- and that the
builder that grabs the ACTIVE sheet (_sheet_lists) never lands on his.

2026-10-12: he added his own cells above the table (a Kalshi price ->
sportsbook odds converter in K6, total staked in M5) and asked that K6 be
kept, and for charts beside the table that he can delete. The charts are the
one thing the model draws on his sheet; the tests at the bottom pin that they
go on once, that a deleted one stays deleted, and that none of it touches a
cell.
"""
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from nflmodel import excel
from nflmodel.betlog import (TRACKER_LAST_ROW, TRACKER_SHEET, load_tracker,
                             tracker_fingerprint)

KALSHI = "=IF(J6>=0.5, -ROUND(J6/(1-J6)*100,0), ROUND((1-J6)/J6*100,0))"
MODEL_CHARTS = ["Profit since the start", "Winning weeks", "Your win %", "Made money"]


def _his_workbook(path):
    """A weekly workbook as Excel leaves it after he has typed into it."""
    wb = Workbook()
    wb.active.title = "Lists"
    wb.create_sheet("This Week's Bets")["A1"] = "old model sheet"
    ws = wb.create_sheet(TRACKER_SHEET)
    ws["A1"] = "BET TRACKER"
    ws.merge_cells("A1:P1")
    ws["B4"] = 1000
    ws["A11"] = "Date"
    ws["A12"] = datetime(2026, 9, 27)
    ws["B12"] = 3
    ws["C12"] = "TEN @ NY "
    ws["D12"] = "MONEYLINE"
    ws["E12"] = "HOME"
    ws["G12"] = -127
    ws["H12"] = 5
    ws["I12"] = "W"
    ws["J12"] = '=IF($I12="W",IF($G12<0,$H12*100/ABS($G12),$H12*$G12/100),"")'
    ws["Q12"] = "his own note"                         # past the last column
    ws["J5"], ws["K5"] = "Kalshi Odds ", "Sports Book Odds"      # his, 2026-10-12
    ws["J6"] = 0.53
    ws["K6"] = KALSHI
    ws["M4"], ws["M5"] = "Total Money Bet ", "=SUM(H12:H300)"
    ws["C12"].fill = PatternFill("solid", fgColor="FFFF00")
    dv = DataValidation(type="list", formula1='"W,L,Push"', allow_blank=True)
    dv.sqref = "I12:I500"
    ws.add_data_validation(dv)
    wb.create_sheet("Power Ratings")["A1"] = "old ratings"
    wb.create_sheet("My Notes")["A1"] = "a sheet he added himself"
    wb.create_sheet("Model Data").sheet_state = "hidden"
    wb.save(str(path))
    return path


def _cells(ws):
    return {(c.row, c.column): c.value
            for row in ws.iter_rows() for c in row if c.value not in (None, "")}


def test_tracker_is_carried_cell_for_cell(tmp_path):
    out = tmp_path / "output"
    out.mkdir()
    src = _his_workbook(out / "NFL_Model_2026_Week04.xlsx")
    tracker = load_tracker(out, tmp_path / "bet_log.csv")
    assert tracker.kind == "workbook"

    wb = excel._start_workbook(tracker)
    # Only his sheets survive, plus a new Lists that is ACTIVE.
    assert wb.sheetnames == ["Lists", TRACKER_SHEET, "My Notes"]
    assert wb.active.title == "Lists"

    excel._sheet_lists(wb, [(3.0, 0.08)])        # takes over the active sheet
    assert TRACKER_SHEET in wb.sheetnames, "_sheet_lists renamed his tracker"
    wb.create_sheet("This Week's Bets")["A1"] = "new model sheet"
    wb.create_sheet("Power Ratings")
    wb.create_sheet("Bet Log")
    excel._place_tracker(wb)
    assert wb.sheetnames == ["Lists", "This Week's Bets", "Power Ratings",
                             TRACKER_SHEET, "Bet Log", "My Notes"]
    dest = out / "NFL_Model_2026_Week05.xlsx"
    wb.save(str(dest))

    before = load_workbook(str(src))[TRACKER_SHEET]
    after_wb = load_workbook(str(dest))
    after = after_wb[TRACKER_SHEET]
    assert _cells(after) == _cells(before)
    assert tracker_fingerprint(dest) == tracker_fingerprint(src)
    assert after["Q12"].value == "his own note"
    assert after["J12"].value.startswith("=IF(")
    assert after["K6"].value == KALSHI, "his Kalshi converter was not carried"
    assert after["J6"].value == 0.53
    assert after["M5"].value == "=SUM(H12:H300)"
    assert after["C12"].fill.fgColor.rgb.endswith("FFFF00")
    assert "A1:P1" in {str(r) for r in after.merged_cells.ranges}
    assert [str(d.sqref) for d in after.data_validations.dataValidation] == ["I12:I500"]
    # The old model sheets are gone, replaced by the new ones.
    assert after_wb["This Week's Bets"]["A1"].value == "new model sheet"
    assert "Model Data" not in after_wb.sheetnames
    assert after_wb["My Notes"]["A1"].value == "a sheet he added himself"


def test_every_sheet_the_builders_create_is_registered():
    # A model sheet missing from MODEL_SHEETS would be treated as HIS and
    # carried forward stale every week, next to its rebuilt twin.
    import re
    src = open(excel.__file__).read()
    created = set(re.findall(r'create_sheet\("([^"]+)"', src))
    created.discard("Bet Tracker")
    assert created <= excel.MODEL_SHEETS, created - excel.MODEL_SHEETS


def test_fresh_workbook_when_he_has_no_tracker(tmp_path):
    out = tmp_path / "output"
    out.mkdir()
    tracker = load_tracker(out, tmp_path / "bet_log.csv")
    assert tracker.kind == "none"
    wb = excel._start_workbook(tracker)
    assert wb.sheetnames == ["Sheet"]            # the builders take it from here


# ── the charts beside the table ──────────────────────────────────────────

def _names(ws):
    return [excel._chart_identity(c)[0] for c in ws._charts]


def _carry(out, tmp_path):
    return excel._start_workbook(load_tracker(out, tmp_path / "bet_log.csv", announce=False))


def test_charts_go_on_once_and_a_deleted_one_stays_deleted(tmp_path, capsys):
    out = tmp_path / "output"
    out.mkdir()
    _his_workbook(out / "NFL_Model_2026_Week04.xlsx")       # built before the charts existed

    wb = _carry(out, tmp_path)
    ws = wb[TRACKER_SHEET]
    assert _names(ws) == MODEL_CHARTS
    excel._sheet_chart_data(wb)

    # In Excel he deletes one, drags another and makes a chart of his own.
    ws._charts = [c for c in ws._charts if excel._chart_identity(c)[0] != "Made money"]
    ws._charts[0].anchor = "T40"
    own = BarChart()
    own.add_data(Reference(ws, min_col=8, min_row=11, max_row=20), titles_from_data=True)
    ws.add_chart(own, "AL12")
    week5 = out / "NFL_Model_2026_Week05.xlsx"
    wb.save(str(week5))
    newer = (out / "NFL_Model_2026_Week04.xlsx").stat().st_mtime + 60
    os.utime(week5, (newer, newer))

    capsys.readouterr()
    after = _carry(out, tmp_path)[TRACKER_SHEET]
    said = capsys.readouterr().out
    assert "3 redrawn" in said and "left off: profit by bet type" in said, \
        "the weekly run must say which charts it refreshed"
    by_name = {excel._chart_identity(c)[0]: c for c in after._charts}
    assert "Made money" not in by_name, "a chart he deleted was put back"
    assert set(MODEL_CHARTS) - {"Made money"} <= set(by_name)
    moved = by_name["Profit since the start"].anchor._from
    assert (moved.col, moved.row) == (19, 39), "a chart he moved went back to its old spot"
    assert sum(1 for c in after._charts if not excel._chart_identity(c)[1]) == 1, "his own chart was lost"
    assert after["K6"].value == KALSHI


def test_every_chart_series_has_data_behind_it():
    wb = Workbook()
    ws = wb.create_sheet(TRACKER_SHEET)
    excel._tracker_charts(ws, install_all=True)
    excel._sheet_chart_data(wb)
    data = wb["Chart Data"]
    assert data.sheet_state == "hidden"
    checked = 0
    for chart in ws._charts:
        for series in chart.series:
            for src in (series.val, series.yVal, series.xVal, series.cat):
                if src is None:
                    continue
                sheet, rng = (src.numRef or src.strRef).f.split("!")
                assert sheet.strip("'") == "Chart Data"
                empty = [c.coordinate for row in data[rng.replace("$", "")] for c in row if c.value is None]
                assert not empty, f"{rng} plots empty cells: {empty[:5]}"
                checked += 1
        title = chart.title.tx.strRef
        if title is not None:                     # the bankroll chart's live title
            assert data[title.f.split("!")[1].replace("$", "")].value.startswith('="BANKROLL $"')
    assert checked == 14              # 2 + 4 + 4 + 4 ranges across the four charts
    # The per-bet rows reach the last row he can type a bet into.
    assert f"$J{TRACKER_LAST_ROW})" in data.cell(row=excel._CD_LAST, column=1).value
    assert data.cell(row=excel._CD_LAST + 1, column=1).value is None


def test_a_tracker_rebuilt_from_scratch_keeps_his_converter():
    # Lost output/: the sheet is rebuilt and his rows come back from the CSV.
    wb = Workbook()
    excel._sheet_tracker(wb, [])
    ws = wb[TRACKER_SHEET]
    assert ws["K6"].value == KALSHI == excel.KALSHI_TO_SPORTSBOOK
    assert (ws["J5"].value, ws["K5"].value, ws["J6"].value) == ("Kalshi Odds", "Sports Book Odds", 0.53)
    assert ws["M5"].value == "=SUM($H$12:$H$500)"
    assert _names(ws) == MODEL_CHARTS
