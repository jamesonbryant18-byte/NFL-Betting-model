"""
Tests for betlog.py — the Bet Tracker is Jameson's, read and never edited.

The scenario these pin happened (2026-09-30): he deleted rows in his newest
workbook and the old union of every copy put them back from older copies.
Now the newest workbook's tracker IS the tracker; older copies and the CSV
never add a row to it, and a copy he edited is backed up, not overwritten.

Run: .venv/bin/python -W ignore -m pytest tests/test_betlog.py -q
"""
import os
import sys
import time
from datetime import date, datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))

import pytest
from openpyxl import Workbook

from nflmodel import betlog
from nflmodel.betlog import (BET_LOG_COLUMNS, bet_key, clv_for_row,
                             load_tracker, normalize_matchup, parse_row,
                             protect_tracker_copies, read_csv_rows,
                             read_tracker_rows, record_tracker_state,
                             tracker_fingerprint, write_csv_rows)


# ── fixtures ───────────────────────────────────────────────────

def _write_tracker(path, rows, sheet="Bet Tracker"):
    """A workbook shaped like excel._sheet_tracker's: data from row 12,
    A..I hand-entered, L/M/P extras. Formula columns are left empty."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.cell(row=11, column=1, value="Date")
    for i, rec in enumerate(rows):
        r = 12 + i
        for c, v in enumerate(rec[:9], start=1):
            ws.cell(row=r, column=c, value=v)
        for off, col in enumerate((12, 13, 16)):
            idx = 9 + off
            if len(rec) > idx and rec[idx] is not None:
                ws.cell(row=r, column=col, value=rec[idx])
    wb.save(str(path))
    return path


def _set_mtime(path, when: float):
    os.utime(str(path), (when, when))


# The same bet, at different points in its life.
SHARED_OLD = ["9/13/2026", 1, "MIA @ LV", "SPREAD", "MIA +3.5", 3.5, -110, 50, None,
              None, None, None]
SHARED_NEW = ["9/13/2026", 1, "MIA @ LV", "SPREAD", "MIA +3.5", 3.5, -110, 75, "W",
              1.5, -108, 2.1]
OLD_ONLY = ["9/14/2026", 1, "DEN @ KC", "MONEYLINE", "DEN +140", None, 140, 25, "L",
            None, 150, 0.04]
NEW_ONLY = ["9/20/2026", 2, "KC @ NYG", "SPREAD", "NYG +6.5", 6.5, -105, 40, None,
            None, None, None]
CSV_ONLY = ["9/13/2026", "1", "DAL @ NYG", "SPREAD", "NYG +3.0", "3.0", "-115", "50",
            "", "", "", ""]


@pytest.fixture
def stores(tmp_path):
    """Two workbooks and a CSV with controlled mtimes.

    newest:  Week02 workbook  — edited stake/result on the shared bet, + NEW_ONLY
    middle:  CSV mirror       — the shared bet as first logged, + CSV_ONLY
    oldest:  Week01 workbook  — the shared bet as first logged, + OLD_ONLY
    """
    out = tmp_path / "output"
    out.mkdir()
    wk1 = _write_tracker(out / "NFL_Model_2026_Week01.xlsx", [SHARED_OLD, OLD_ONLY])
    wk2 = _write_tracker(out / "NFL_Model_2026_Week02.xlsx", [SHARED_NEW, NEW_ONLY])
    csv_path = tmp_path / "data" / "bet_log.csv"
    write_csv_rows([SHARED_OLD, CSV_ONLY], csv_path)

    now = time.time()
    _set_mtime(wk1, now - 3 * 3600)
    _set_mtime(csv_path, now - 2 * 3600)
    _set_mtime(wk2, now - 1 * 3600)
    return {"out": out, "wk1": wk1, "wk2": wk2, "csv": csv_path}


# ── reading the stores ──────────────────────────────────────────

def test_read_tracker_rows_matches_sheet_layout(stores):
    rows = read_tracker_rows(stores["wk2"])
    assert len(rows) == 2
    assert all(len(r) == 12 for r in rows)
    shared = rows[0]
    assert shared[:9] == SHARED_NEW[:9]
    # L, M, P land in slots 9, 10, 11 — the same order the CSV uses.
    assert shared[9:] == [1.5, -108, 2.1]


def test_read_tracker_rows_is_empty_for_missing_file_or_sheet(tmp_path):
    assert read_tracker_rows(tmp_path / "nope.xlsx") == []
    other = _write_tracker(tmp_path / "other.xlsx", [SHARED_OLD], sheet="Not It")
    assert read_tracker_rows(other) == []
    junk = tmp_path / "junk.xlsx"
    junk.write_bytes(b"this is not a workbook")
    assert read_tracker_rows(junk) == []


def test_csv_roundtrip_pads_and_truncates(tmp_path):
    p = tmp_path / "log.csv"
    write_csv_rows([SHARED_OLD[:5], SHARED_NEW + ["extra"]], p)
    rows = read_csv_rows(p)
    assert len(rows) == 2 and all(len(r) == 12 for r in rows)
    assert rows[0][:5] == [str(v) for v in SHARED_OLD[:5]]
    assert rows[0][5:] == [""] * 7
    assert p.read_text().splitlines()[0] == ",".join(BET_LOG_COLUMNS)
    assert read_csv_rows(tmp_path / "missing.csv") == []


# ── identity ────────────────────────────────────────────────────

def test_normalize_matchup_only_touches_separator_case_and_spacing():
    for s in ("MIA @ LV", "MIA@LV", "mia at lv", "  mia   @  lv "):
        assert normalize_matchup(s) == "MIA @ LV"
    assert normalize_matchup("LV vs MIA") == "LV VS MIA"
    assert normalize_matchup(None) == ""
    # 'AT' must be the separator, not part of a team code (ATL).
    assert normalize_matchup("ATL @ PIT") == "ATL @ PIT"
    assert normalize_matchup("atl at pit") == "ATL @ PIT"


def test_bet_key_ignores_the_editable_fields_and_the_date_type():
    assert bet_key(SHARED_OLD) == bet_key(SHARED_NEW)
    typed = list(SHARED_OLD)
    typed[0] = datetime(2026, 9, 13)              # what openpyxl hands back
    typed[2], typed[3], typed[4] = "mia@lv", "spread", "mia  +3.5"
    assert bet_key(typed) == bet_key(SHARED_OLD)
    assert bet_key(SHARED_OLD) != bet_key(OLD_ONLY)


# ── one tracker, read only ─────────────────────────────────────

def test_load_tracker_reads_only_the_newest_workbook(stores):
    t = load_tracker(stores["out"], stores["csv"])
    assert t.kind == "workbook" and t.path == stores["wk2"]
    # Exactly what he left in the newest copy: nothing resurrected from the
    # older workbook or the CSV, even though they hold rows it does not.
    assert [r[:9] for r in t.rows] == [SHARED_NEW[:9], NEW_ONLY[:9]]
    assert t.fingerprint == tracker_fingerprint(stores["wk2"])


def test_a_row_he_deleted_does_not_come_back(stores):
    # He deletes NEW_ONLY and fixes nothing else; the older copies still hold
    # rows the newest one lacks. The tracker is what he left: one row.
    _write_tracker(stores["wk2"], [SHARED_NEW])
    _set_mtime(stores["wk2"], time.time())
    t = load_tracker(stores["out"], stores["csv"])
    assert [r[:9] for r in t.rows] == [SHARED_NEW[:9]]


def test_csv_backup_is_read_only_when_no_workbook_exists(tmp_path):
    out = tmp_path / "output"
    out.mkdir()
    csv_path = tmp_path / "bet_log.csv"
    write_csv_rows([CSV_ONLY], csv_path)
    t = load_tracker(out, csv_path)
    assert t.kind == "csv" and len(t.rows) == 1
    assert load_tracker(tmp_path / "missing", tmp_path / "none.csv").kind == "none"


def test_unedited_copies_raise_no_warning(stores):
    t = load_tracker(stores["out"], stores["csv"])
    record_tracker_state(stores["out"], t, stores["wk2"])
    record_tracker_state(stores["out"],
                         load_tracker(stores["out"], stores["csv"]), stores["wk2"])
    # Week01's tracker was never recorded as edited; overwriting a file that
    # is not there is nothing to protect.
    assert protect_tracker_copies(t, stores["out"] / "NFL_Model_2026_Week03.xlsx",
                                  stores["out"]) == []


def test_an_older_copy_edited_after_it_was_carried_is_reported(stores):
    t = load_tracker(stores["out"], stores["csv"])
    # The model has seen Week01 as it is now...
    state = {"known": {"NFL_Model_2026_Week01.xlsx": tracker_fingerprint(stores["wk1"])}}
    (stores["out"] / ".tracker_state.json").write_text(__import__("json").dumps(state))
    # ...then he types a new bet into Week01, but Week02 stays the newest save.
    late = ["9/21/2026", 2, "BUF @ MIA", "SPREAD", "MIA +2.5", 2.5, -110, 5, None,
            None, None, None]
    _write_tracker(stores["wk1"], [SHARED_OLD, OLD_ONLY, late])
    _set_mtime(stores["wk1"], time.time() - 1800)
    _set_mtime(stores["wk2"], time.time())
    t = load_tracker(stores["out"], stores["csv"])
    warns = protect_tracker_copies(t, stores["out"] / "NFL_Model_2026_Week03.xlsx",
                                   stores["out"])
    assert len(warns) == 1 and "NFL_Model_2026_Week01.xlsx" in warns[0]
    assert "BUF @ MIA" in warns[0] and "nothing was merged" in warns[0]
    # The carried tracker did not change.
    assert [r[:9] for r in t.rows] == [SHARED_NEW[:9], NEW_ONLY[:9]]


def test_an_edited_copy_is_backed_up_before_it_is_overwritten(stores):
    # Re-running Week 1 while Week 2 is the newest save: Week01 would be
    # rebuilt around Week02's tracker. Its own differs and the model has no
    # record of it, so it is copied to output/backup/ first.
    t = load_tracker(stores["out"], stores["csv"])
    before = stores["wk1"].read_bytes()
    warns = protect_tracker_copies(t, stores["wk1"], stores["out"])
    backups = list((stores["out"] / "backup").glob("NFL_Model_2026_Week01.*.xlsx"))
    assert len(backups) == 1 and backups[0].read_bytes() == before
    assert len(warns) == 1 and "backup" in warns[0]
    assert stores["wk1"].read_bytes() == before      # protect never writes it


# ── typed view ──────────────────────────────────────────────────

@pytest.mark.parametrize("raw", [
    "9/13/2026", "09/13/2026", "2026-09-13", "2026-09-13 00:00:00",
    datetime(2026, 9, 13, 0, 0), date(2026, 9, 13),
])
def test_parse_row_accepts_every_date_form(raw):
    row = list(SHARED_NEW)
    row[0] = raw
    assert parse_row(row)["date"] == date(2026, 9, 13)


def test_parse_row_types_and_blanks():
    p = parse_row(CSV_ONLY)
    assert p["date"] == date(2026, 9, 13)
    assert p["week"] == 1
    assert p["matchup"] == "DAL @ NYG" and p["market"] == "SPREAD"
    assert p["line_taken"] == 3.0 and p["odds"] == -115.0 and p["stake"] == 50.0
    assert p["result"] == ""
    assert p["closing_line"] is None and p["closing_odds"] is None
    assert p["model_edge"] is None

    p = parse_row(SHARED_NEW)
    assert p["result"] == "W" and p["closing_line"] == 1.5
    assert parse_row(["", "", "", "", "", "", "", "", "Push"])["result"] == "PUSH"
    assert parse_row(["not a date", "wk 2"])["date"] is None
    assert parse_row(["not a date", "wk 2"])["week"] == 2


# ── CLV ─────────────────────────────────────────────────────────

FAKE_PTS = [(0.0, 0.01), (1.5, 0.02), (2.5, 0.03), (3.5, 0.081), (7.5, 0.06)]


def test_clv_spread_crossing_three_uses_the_line_taken_rate():
    parsed = parse_row(["9/13/2026", 1, "MIA @ LV", "SPREAD", "MIA +3.5",
                        3.5, -110, 50, "", 2.5, -110, ""])
    pts, prob = clv_for_row(parsed, FAKE_PTS)
    assert pts == pytest.approx(1.0)
    # One point gained at a taken line of 3.5 is worth the crossing-3 rate.
    assert prob == pytest.approx(0.081)

    # Negative CLV: took -3.5, closed -2.5 (both from the bettor's side).
    parsed = parse_row(["9/13/2026", 1, "MIA @ LV", "SPREAD", "LV -3.5",
                        -3.5, -110, 50, "", -2.5, -110, ""])
    pts, prob = clv_for_row(parsed, FAKE_PTS)
    assert pts == pytest.approx(-1.0)
    assert prob == pytest.approx(-0.081)      # lookup is on |line|


def test_clv_falls_back_to_three_percent_and_handles_missing():
    parsed = parse_row(["", "", "X @ Y", "SPREAD", "X +4", 4.0, -110, 10, "", 2.0])
    assert clv_for_row(parsed, FAKE_PTS) == (pytest.approx(2.0), pytest.approx(0.06))
    assert clv_for_row(parsed, None) == (pytest.approx(2.0), pytest.approx(0.06))
    no_close = parse_row(["", "", "X @ Y", "SPREAD", "X +4", 4.0, -110, 10])
    assert clv_for_row(no_close, FAKE_PTS) == (None, None)


def test_clv_moneyline_is_the_price_difference():
    parsed = parse_row(["", "", "X @ Y", "MONEYLINE", "X +140", "", 140, 10, "",
                        "", 120, ""])
    pts, prob = clv_for_row(parsed, FAKE_PTS)
    assert pts is None
    # +140 -> 41.67%, +120 -> 45.45%: the close was shorter, CLV positive.
    assert prob == pytest.approx(100 / 220 - 100 / 240)
    assert clv_for_row(parse_row(["", "", "X @ Y", "MONEYLINE", "X", "", 140]),
                       FAKE_PTS) == (None, None)


def test_home_away_side_resolves_to_team():
    from nflmodel.betlog import parse_row
    home = parse_row(["2026-09-20", 2, "CAR @ ATL", "SPREAD", "HOME", 2.5, "", 5])
    away = parse_row(["2026-09-17", 2, "DET @ BUF", "MONEYLINE", "away", "", "", 5])
    team = parse_row(["2026-09-20", 2, "CLE @ TB", "SPREAD", "CLE +8.5", "", "", 5])
    assert home["bet_side"] == "ATL"
    assert away["bet_side"] == "DET"
    assert team["bet_side"] == "CLE +8.5"


def test_an_edit_to_an_older_copy_is_reported_once(stores):
    t = load_tracker(stores["out"], stores["csv"])
    record_tracker_state(stores["out"], t, stores["wk2"])     # baselines Week01 too
    late = ["9/21/2026", 2, "BUF @ MIA", "SPREAD", "MIA +2.5", 2.5, -110, 5, None,
            None, None, None]
    _write_tracker(stores["wk1"], [SHARED_OLD, OLD_ONLY, late])
    _set_mtime(stores["wk1"], time.time() - 1800)
    _set_mtime(stores["wk2"], time.time())
    target = stores["out"] / "NFL_Model_2026_Week03.xlsx"
    t = load_tracker(stores["out"], stores["csv"])
    first = protect_tracker_copies(t, target, stores["out"])
    assert len(first) == 1 and "BUF @ MIA" in first[0]
    assert protect_tracker_copies(t, target, stores["out"]) == []
